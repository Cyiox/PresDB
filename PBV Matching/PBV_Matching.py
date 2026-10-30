import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from fuzzywuzzy import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
FORECAST_DIR = PROJECT_ROOT / "Forecast Data" / "Current"
PBV_PATH = SCRIPT_DIR / "PBVExtractForCEDAC.xlsx"
MATCH_OUTPUT = PROJECT_ROOT / "PBV Matching" / "PBV_Matching3.xlsx"
SUMMARY_OUTPUT = PROJECT_ROOT / "PBV Matching" / "2030_Forecast_PBV_Match_Summary.xlsx"
UNMATCHED_OUTPUT = PROJECT_ROOT / "PBV Matching" / "Unmatched_PBV_Projects.xlsx"


def resolve_latest_forecast_path() -> Path:
    candidates = list(FORECAST_DIR.glob("2030 Forecast Master*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No 2030 forecast workbook found in {FORECAST_DIR}")

    def sort_key(path: Path) -> tuple[datetime, float, str]:
        match = re.search(r"Master\s+(\d{1,2})-(\d{1,2})(?:-(\d{2,4}))?", path.stem)
        if not match:
            return (datetime.min, path.stat().st_mtime, path.name.lower())
        year_text = match.group(3)
        year = datetime.now().year if year_text is None else int(year_text)
        if year < 100:
            year += 2000
        try:
            report_date = datetime(year, int(match.group(1)), int(match.group(2)))
        except ValueError:
            report_date = datetime.min
        return (report_date, path.stat().st_mtime, path.name.lower())

    return max(candidates, key=sort_key)


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    result = df.copy()
    result.columns = [str(column).strip() for column in result.columns]
    return result.fillna("")


def require_columns(df: pd.DataFrame, source_name: str, columns: list[str]) -> None:
    missing = [column for column in columns if column not in df.columns]
    if missing:
        raise ValueError(f"{source_name} is missing required columns: {', '.join(missing)}")


def text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def normalize_name(value: object) -> str:
    result = text(value).lower()
    result = re.sub(r"\b(street|strt)\b", "st", result)
    result = re.sub(r"\b(avenue)\b", "ave", result)
    result = re.sub(r"\b(building)\b", "bldg", result)
    result = re.sub(r"\bveterans\b", "veteran", result)
    result = re.sub(
        r"\b(apartments|apts|apartment|housing|house|residences?|senior|estates?|homes?|village|towers?|place|commons|court|manor|the)\b",
        " ",
        result,
    )
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", result)).strip()


def normalize_city(value: object) -> str:
    result = normalize_name(value)
    if result.startswith("boston "):
        return "boston"
    return result


def normalize_address(value: object) -> str:
    result = text(value).lower()
    replacements = {
        "street": "st", "strt": "st", "road": "rd", "avenue": "ave",
        "drive": "dr", "lane": "ln", "court": "ct", "boulevard": "blvd",
    }
    for source, target in replacements.items():
        result = re.sub(rf"\b{source}\b", target, result)
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", result)).strip()


def address_number(value: object) -> int | None:
    match = re.match(r"^\s*(\d+)", text(value))
    return int(match.group(1)) if match else None


def address_street_name(value: object) -> str:
    result = re.sub(r"^\s*\d+\s*-?\s*\d*\s*", "", text(value).lower())
    result = re.sub(
        r"\b(street|st|avenue|ave|road|rd|drive|dr|lane|ln|court|ct|circle|cir|way|boulevard|blvd|place|pl)\b\.?$",
        "",
        result,
    )
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", result)).strip()


def address_score(forecast_address: object, pbv_address: object) -> tuple[int, list[str]]:
    forecast = normalize_address(forecast_address)
    pbv = normalize_address(pbv_address)
    if not forecast or not pbv:
        return 0, ["MISSING_ADDRESS"]

    flags: list[str] = []
    forecast_number = address_number(forecast_address)
    pbv_number = address_number(pbv_address)
    if forecast_number is not None and pbv_number is not None:
        if abs(forecast_number - pbv_number) > 20:
            return 0, ["ADDRESS_NUMBER_MISMATCH"]
        if forecast_number % 2 != pbv_number % 2:
            flags.append("ODD_EVEN_MISMATCH")

    street_score = fuzz.ratio(address_street_name(forecast_address), address_street_name(pbv_address))
    score = fuzz.ratio(forecast, pbv)
    if street_score < 90 or score < 80:
        flags.append("ADDRESS_MISMATCH")
    return score, flags


def unit_evidence(forecast_units: object, pbv_units: object) -> tuple[int, str, list[str]]:
    try:
        forecast_number = float(str(forecast_units).replace(",", ""))
        pbv_number = float(str(pbv_units).replace(",", ""))
    except (TypeError, ValueError):
        return 0, "Unit count missing or invalid", ["MISSING_UNITS"]

    difference = abs(forecast_number - pbv_number)
    if difference == 0:
        return 5, f"Exact unit match ({int(forecast_number)})", []
    if difference <= 2:
        return 3, f"Units very close (difference: {int(difference)})", []
    if difference <= 10:
        return -5, f"Units differ moderately (difference: {int(difference)})", ["UNIT_MISMATCH_MODERATE"]
    if difference <= 30:
        return -10, f"Units differ significantly (difference: {int(difference)})", ["UNIT_MISMATCH_LARGE"]
    return -15, f"Units differ greatly (difference: {int(difference)})", ["UNIT_MISMATCH_VERY_LARGE"]


def candidate_for_pair(forecast_row: pd.Series, pbv_row: pd.Series) -> dict[str, object] | None:
    if normalize_city(forecast_row["City"]) != normalize_city(pbv_row["City/Town"]):
        return None

    names = [normalize_name(forecast_row["PropertyName"])]
    aka = text(forecast_row.get("Project/AKA", ""))
    names.extend(normalize_name(value) for value in re.split(r"[,/|;]", aka))
    names = [name for name in names if name]
    pbv_name = normalize_name(pbv_row["Project"])
    if not names or not pbv_name:
        return None

    name_score = max(fuzz.ratio(name, pbv_name) for name in names)
    addr_score, address_flags = address_score(forecast_row["Street"], pbv_row["Address"])
    if name_score < 85 and addr_score < 80:
        return None

    flags = list(address_flags)
    unit_adjustment, unit_note, unit_flags = unit_evidence(
        forecast_row.get("TotalUnits", ""), pbv_row.get("Total Units", "")
    )
    flags.extend(unit_flags)
    if name_score >= 95 and addr_score >= 80:
        method = "Exact/Strong Name + Address Match"
        confidence = max(name_score, addr_score) + unit_adjustment
    elif name_score >= 85:
        method = "Fuzzy Name Match"
        confidence = name_score + unit_adjustment
        if addr_score < 80:
            flags.append("ADDRESS_MISMATCH")
    else:
        method = "Street Address Match"
        confidence = addr_score + unit_adjustment
        flags.append("NAME_MISMATCH")

    return {
        "forecast_idx": forecast_row.name,
        "pbv_idx": pbv_row.name,
        "match_method": method,
        "confidence": max(40, min(100, confidence)),
        "base_confidence": max(name_score, addr_score),
        "name_similarity": name_score,
        "address_similarity": addr_score,
        "flags": sorted(set(flags)),
        "unit_note": unit_note,
    }


def match_properties(df_forecast: pd.DataFrame, df_pbv: pd.DataFrame) -> list[dict[str, object]]:
    candidates: list[dict[str, object]] = []
    for _, forecast_row in df_forecast.iterrows():
        for _, pbv_row in df_pbv.iterrows():
            candidate = candidate_for_pair(forecast_row, pbv_row)
            if candidate is not None:
                candidates.append(candidate)

    matches: list[dict[str, object]] = []
    used_forecast: set[object] = set()
    used_pbv: set[object] = set()
    for candidate in sorted(candidates, key=lambda item: (item["confidence"], item["base_confidence"]), reverse=True):
        if candidate["forecast_idx"] in used_forecast or candidate["pbv_idx"] in used_pbv:
            continue
        matches.append(candidate)
        used_forecast.add(candidate["forecast_idx"])
        used_pbv.add(candidate["pbv_idx"])
    return matches


def merged_output(df_forecast: pd.DataFrame, df_pbv: pd.DataFrame, matches: list[dict[str, object]]) -> pd.DataFrame:
    rows = []
    for match in matches:
        forecast_row = df_forecast.loc[match["forecast_idx"]].to_dict()
        pbv_row = df_pbv.loc[match["pbv_idx"]].to_dict()
        merged = dict(forecast_row)
        merged.update({f"PBV_{key}": value for key, value in pbv_row.items()})
        merged.update(
            {
                "Match_Method": match["match_method"],
                "Match_Confidence": match["confidence"],
                "Base_Confidence": match["base_confidence"],
                "Name_Similarity": match["name_similarity"],
                "Address_Similarity": match["address_similarity"],
                "Flags": "|".join(match["flags"]),
                "Unit_Note": match["unit_note"],
            }
        )
        rows.append(merged)
    return pd.DataFrame(rows).sort_values("Match_Confidence", ascending=True)


def highlight_address_mismatches(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = {cell.value: cell.column for cell in sheet[1]}
    flags_column = headers.get("Flags")
    street_column = headers.get("Street")
    address_column = headers.get("Address") or headers.get("PBV_Address")
    if not flags_column or not street_column or not address_column:
        return
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row=row, column=flags_column).value
        if value and "ADDRESS_MISMATCH" in str(value):
            sheet.cell(row=row, column=street_column).fill = yellow_fill
            sheet.cell(row=row, column=address_column).fill = yellow_fill
    workbook.save(path)


def main() -> None:
    forecast_path = resolve_latest_forecast_path()
    if not PBV_PATH.exists():
        raise FileNotFoundError(f"PBV workbook not found: {PBV_PATH}")

    df_forecast = normalize_columns(pd.read_excel(forecast_path, sheet_name=0))
    df_pbv = normalize_columns(pd.read_excel(PBV_PATH, sheet_name=0))
    require_columns(df_forecast, forecast_path.name, ["PropertyName", "City", "Street", "TotalUnits"])
    require_columns(df_pbv, PBV_PATH.name, ["Project", "City/Town", "Address", "Total Units"])

    print("=== PBV MATCHING ===")
    print(f"Forecast workbook: {forecast_path.name} (first worksheet)")
    print(f"Forecast properties: {len(df_forecast)}")
    print(f"PBV projects: {len(df_pbv)}")
    matches = match_properties(df_forecast, df_pbv)
    print(f"Matches found: {len(matches)}")

    if matches:
        merged = merged_output(df_forecast, df_pbv, matches)
        merged.to_excel(MATCH_OUTPUT, index=False)
        summary_columns = [
            "PropertyID", "PropertyName", "Project/AKA", "City", "Street", "TotalUnits",
            "PBV_Project", "PBV_City/Town", "PBV_Address", "PBV_Total Units", "PBV_Contract Number",
            "PBV_RAA #", "PBV_Status:", "Match_Method", "Match_Confidence", "Base_Confidence",
            "Name_Similarity", "Address_Similarity", "Flags", "Unit_Note",
        ]
        summary = merged[[column for column in summary_columns if column in merged.columns]]
        summary.to_excel(SUMMARY_OUTPUT, index=False)
        highlight_address_mismatches(SUMMARY_OUTPUT)
        print(f"Saved full matches: {MATCH_OUTPUT.name}")
        print(f"Saved review summary: {SUMMARY_OUTPUT.name}")

    matched_pbv = {match["pbv_idx"] for match in matches}
    unmatched_pbv = df_pbv[~df_pbv.index.isin(matched_pbv)]
    unmatched_pbv.to_excel(UNMATCHED_OUTPUT, index=False)
    print(f"Unmatched PBV projects: {len(unmatched_pbv)}")
    print(f"Saved unmatched projects: {UNMATCHED_OUTPUT.name}")


if __name__ == "__main__":
    main()